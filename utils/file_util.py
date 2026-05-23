"""
파일/폴더 처리 유틸리티
- find_master_excel   : input/ 폴더에서 마스터 엑셀 찾기
- read_master_excel   : 마스터 엑셀 행 목록 반환
- find_pdf_for_row    : 이름_유저번호1_유저번호2.pdf 파일 탐색 (flat 구조)
- write_bigo_to_excel : 비고 열(I열, column=9)에 결과 기록
- parse_rrn           : 주민번호 문자열 → (앞6, 뒤7) [레거시]
"""

from __future__ import annotations

import re
import unicodedata
from pathlib import Path
from datetime import datetime

import openpyxl

# 비고 열 위치 (I열, 1-indexed)
BIGO_COL = 9


# ── 마스터 엑셀 탐색 ─────────────────────────────────────────────────────────

def find_master_excel(input_dir: Path) -> Path | None:
    """input/ 폴더에서 .xlsx/.xls 파일 하나 반환. 여러 개면 첫 번째."""
    excels = list(input_dir.glob("*.xlsx")) + list(input_dir.glob("*.xls"))
    excels = [e for e in excels if not e.name.startswith("~")]  # 임시파일 제외
    if not excels:
        print(f"[파일] ⚠️ 마스터 엑셀 없음: {input_dir}")
        return None
    if len(excels) > 1:
        print(f"[파일] ⚠️ 마스터 엑셀 여러 개: {[e.name for e in excels]} → 첫 번째 사용")
    return excels[0]


# ── 셀 값 변환 헬퍼 ──────────────────────────────────────────────────────────

def _cell_to_str(val, zero_pad: int = 0) -> str:
    """
    셀 값을 문자열로 변환.
    숫자로 저장된 값(020321 → 20321)을 zero_pad 자리로 복원.
    """
    if val is None:
        return ""
    try:
        s = str(int(float(str(val))))
        return s.zfill(zero_pad) if zero_pad else s
    except (ValueError, TypeError):
        return str(val).strip()


# ── 마스터 엑셀 읽기 ─────────────────────────────────────────────────────────

def read_master_excel(excel_path: Path) -> list[dict]:
    """
    열 구조:
      A=N, B=관리코드, C=이름,
      D=주민번호앞(6자리), E=주민번호뒤(7자리),
      F=유저번호1, G=유저번호2, H=PDF, I=비고
    1행은 헤더, 2행부터 읽음.
    반환: [{"이름", "주민번호앞", "주민번호뒤", "유저번호1", "유저번호2", "행번호"}, ...]
    """
    rows = []
    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        ws = wb.active
        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            # 완전 빈 행 건너뜀
            if all(cell is None for cell in row):
                continue

            name      = _cell_to_str(row[2] if len(row) > 2 else None)              # C
            rrn_front = _cell_to_str(row[3] if len(row) > 3 else None, zero_pad=6)  # D, 6자리
            rrn_back  = _cell_to_str(row[4] if len(row) > 4 else None, zero_pad=7)  # E, 7자리
            user_num1 = _cell_to_str(row[5] if len(row) > 5 else None)              # F
            user_num2 = _cell_to_str(row[6] if len(row) > 6 else None)              # G
            bigo      = _cell_to_str(row[8] if len(row) > 8 else None)              # I

            if not name:
                continue

            rows.append({
                "이름":      name,
                "주민번호앞": rrn_front,
                "주민번호뒤": rrn_back,
                "유저번호1":  user_num1,
                "유저번호2":  user_num2,
                "행번호":    i,
                "비고":      bigo,
            })
        print(f"[엑셀] ✅ 마스터 엑셀 {len(rows)}행 읽음: {excel_path.name}")
    except Exception as e:
        print(f"[엑셀] ❌ 마스터 엑셀 읽기 실패: {excel_path.name} — {e}")
    return rows


# ── PDF 탐색 (flat 구조) ──────────────────────────────────────────────────────

def find_pdf_for_row(pdf_dir: Path, name: str, user_num1: str, user_num2: str) -> Path | None:
    """
    pdf/ 폴더(flat)에서 이름_유저번호1_유저번호2 를 포함하는 .pdf 파일 탐색.
    실제 파일명 예: RAP01198_강지운_166163_1910090_중소기업_취업자_소득세_감면신청서_1.pdf
    """
    pattern = unicodedata.normalize("NFC", f"{name}_{user_num1}_{user_num2}").lower()
    for f in pdf_dir.glob("*.pdf"):
        stem_nfc = unicodedata.normalize("NFC", f.stem).lower()
        if pattern in stem_nfc:
            return f
    return None


# ── 비고 열 기록 ─────────────────────────────────────────────────────────────

def write_bigo_to_excel(excel_path: Path, row_number: int, result_text: str) -> bool:
    """
    마스터 엑셀 I열(비고, column=9)의 row_number 행에 result_text 기록.
    매 호출마다 파일을 열고 저장해 계단식 기록 버그를 방지.
    """
    try:
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active
        ws.cell(row=row_number, column=BIGO_COL).value = result_text
        wb.save(excel_path)
        try:
            wb.close()
        except Exception:
            pass
        print(f"[엑셀] ✅ 비고 기록: 행{row_number} → {result_text!r}")
        return True
    except Exception as e:
        print(f"[엑셀] ❌ 비고 기록 실패 (행{row_number} → {result_text!r}): {e}")
        return False


# ── 주민번호 파싱 (레거시) ───────────────────────────────────────────────────

def parse_rrn(rrn_raw: str) -> tuple[str | None, str | None]:
    """주민번호 문자열을 (앞6자리, 뒤7자리)로 분리 [레거시 — 신규 엑셀에서는 미사용]."""
    cleaned = re.sub(r'-', '', rrn_raw).strip()
    if not cleaned.isdigit() or len(cleaned) != 13:
        return None, None
    return cleaned[:6], cleaned[6:]


# ── 실패 로그 기록 ────────────────────────────────────────────────────────────

def write_failed_log(log_path: Path, entries: list[dict]):
    """실패 내역을 log_path 파일에 추가(append) 기록."""
    if not entries:
        return
    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    lines = [f"\n[{ts}] 실패 내역 ({len(entries)}건)"]
    for e in entries:
        lines.append(
            f"  이름={e.get('이름', '')} | 유저번호1={e.get('유저번호1', '')} | "
            f"유저번호2={e.get('유저번호2', '')} | 오류={e.get('오류', '')}"
        )
    content = "\n".join(lines) + "\n"
    try:
        with open(log_path, "a", encoding="utf-8") as f:
            f.write(content)
    except Exception as e:
        print(f"[파일] ❌ failed_log.txt 기록 실패: {e}")
